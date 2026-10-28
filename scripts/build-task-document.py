"""Build the Bulgarian project report from the implementation and verified evidence.
Run after `mvn test` and generating the print/PDF PNGs in output/screenshots.
"""
from pathlib import Path
import re, textwrap
from PIL import Image, ImageDraw, ImageFont, ImageChops
from docx import Document
from docx.shared import Inches, Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'output'
SHOTS = OUT / 'screenshots'
JAVA = ROOT / 'src/main/java/org/codingburgas/demetra'
MONO = '/opt/anaconda3/lib/python3.12/site-packages/matplotlib/mpl-data/fonts/ttf/DejaVuSansMono.ttf'
SANS = str(ROOT / 'src/main/resources/fonts/DejaVuSans.ttf')

def source_image(file, first, last, name):
    rows = (JAVA/file).read_text().splitlines()
    f = ImageFont.truetype(MONO, 19)
    titlef = ImageFont.truetype(SANS, 22)
    content = []
    for n in range(first, min(last,len(rows))+1):
        line = rows[n-1].expandtabs(4)
        # Wrap only for the image; original file is unchanged.
        parts = textwrap.wrap(line, width=103, replace_whitespace=False, drop_whitespace=False) or ['']
        for i, part in enumerate(parts): content.append((str(n) if i==0 else '',part))
    img = Image.new('RGB', (1400, 92 + 29*len(content)), 'white')
    draw = ImageDraw.Draw(img)
    draw.rectangle((0,0,1399,58), fill='#edf1f6')
    draw.text((25,14), file, font=titlef, fill='#162a43')
    for i,(n,line) in enumerate(content):
        y=76+i*29
        draw.text((15,y), n.rjust(3), font=f, fill='#718096')
        color = '#1f2937'
        if line.strip().startswith('//'): color='#4b8061'
        draw.text((83,y),line,font=f,fill=color)
    img.save(SHOTS/name)

source_image('services/OrderFilterService.java',27,51,'code-orders.png')
source_image('services/ReportService.java',30,54,'code-reports.png')
source_image('controllers/ReportController.java',31,43,'code-pdf-route.png')
source_image('services/ReportPdfService.java',43,63,'code-pdf-pages.png')

# Crops of our generated renderings, not screenshots of a live site.
for src,dst,box in [
    ('orders-filtered-1.png','orders-print-crop.png',(55,66,1080,1160)),
    ('reports-1-1.png','reports-controls-crop.png',(65,75,1080,850)),
    ('report-12-2.png','pdf-events-crop.png',(60,60,1070,370)),
]:
    Image.open(SHOTS/src).crop(box).save(SHOTS/dst)

doc=Document()
sec=doc.sections[0]
sec.page_width=Cm(21); sec.page_height=Cm(29.7)
sec.top_margin=Cm(1.7); sec.bottom_margin=Cm(1.6)
sec.left_margin=Cm(1.8); sec.right_margin=Cm(1.8)
sec.header_distance=Cm(.7); sec.footer_distance=Cm(.7)
for sty in ['Normal','Title','Subtitle','Heading 1','Heading 2','Caption']:
    s=doc.styles[sty]; s.font.name='Arial'; s.font.color.rgb=RGBColor(0,0,0)
    s.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Arial')
doc.styles['Normal'].font.size=Pt(10.5)
doc.styles['Normal'].paragraph_format.space_after=Pt(7)
doc.styles['Normal'].paragraph_format.line_spacing=1.12
doc.styles['Title'].font.size=Pt(25)
doc.styles['Title'].paragraph_format.space_after=Pt(13)
doc.styles['Heading 1'].font.size=Pt(17)
doc.styles['Heading 1'].paragraph_format.space_after=Pt(10)
doc.styles['Heading 2'].font.size=Pt(12)
doc.styles['Heading 2'].paragraph_format.space_before=Pt(10)
doc.styles['Caption'].font.size=Pt(9)
doc.styles['Caption'].font.italic=False
doc.styles['Caption'].font.bold=False
doc.styles['Caption'].paragraph_format.space_after=Pt(8)
header=sec.header.paragraphs[0]
header.text='DEMETRA DARIQ  |  Поръчки и отчети'
header.style='Normal'; header.runs[0].font.size=Pt(8)
footer=sec.footer.paragraphs[0]; footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
r=footer.add_run('16 октомври 2026  •  '); r.font.size=Pt(8)
f=OxmlElement('w:fldSimple'); f.set(qn('w:instr'),'PAGE'); footer._p.append(f)

def p(text, bold=False):
    a=doc.add_paragraph(); r=a.add_run(text); r.bold=bold; return a

def h(text): doc.add_heading(text,1)
def sub(text): doc.add_heading(text,2)
def page(): doc.add_page_break()
def code(text):
    a=doc.add_paragraph(); a.paragraph_format.space_after=Pt(6)
    r=a.add_run(text); r.font.name='DejaVu Sans Mono'; r.font.size=Pt(8.5)
    return a

def picture(name,caption,width=6.7):
    a=doc.add_paragraph(); a.paragraph_format.space_after=Pt(3)
    a.alignment=WD_ALIGN_PARAGRAPH.CENTER
    a.add_run().add_picture(str(SHOTS/name),width=Inches(width))
    a.paragraph_format.keep_with_next=True
    # Accessible description stored in the DOCX drawing.
    a._p.xpath('.//wp:docPr')[0].set('descr',caption)
    doc.add_paragraph(caption,'Caption')

def table(headers, rows, widths):
    t=doc.add_table(rows=1, cols=len(headers)); t.alignment=WD_TABLE_ALIGNMENT.CENTER
    t.autofit=False
    for i,w in enumerate(widths): t.columns[i].width=Cm(w)
    for i,s in enumerate(headers): t.rows[0].cells[i].text=str(s)
    for row in rows:
        cells=t.add_row().cells
        for i,s in enumerate(row): cells[i].text=str(s)
    for ri,row in enumerate(t.rows):
        for ci,c in enumerate(row.cells):
            c.width=Cm(widths[ci]); c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcPr=c._tc.get_or_add_tcPr()
            borders=OxmlElement('w:tcBorders')
            for edge in ['top','left','bottom','right']:
                b=OxmlElement('w:'+edge); b.set(qn('w:val'),'single'); b.set(qn('w:sz'),'4'); b.set(qn('w:color'),'D9D9D9'); borders.append(b)
            tcPr.append(borders)
            margins=OxmlElement('w:tcMar')
            for edge in ['top','left','bottom','right']:
                m=OxmlElement('w:'+edge); m.set(qn('w:w'),'95'); m.set(qn('w:type'),'dxa'); margins.append(m)
            tcPr.append(margins)
            fill='dce5f0' if ri==0 else ('f5f7fa' if ri%2==0 else 'FFFFFF')
            sh=OxmlElement('w:shd'); sh.set(qn('w:fill'),fill); tcPr.append(sh)
            for para in c.paragraphs:
                para.paragraph_format.space_after=Pt(2); para.paragraph_format.line_spacing=1.05
                for r in para.runs: r.font.size=Pt(9); r.bold=ri==0
    rep=OxmlElement('w:tblHeader'); t.rows[0]._tr.get_or_add_trPr().append(rep)
    doc.add_paragraph().paragraph_format.space_after=Pt(2)
    return t

# 1
h0=doc.add_paragraph('Поръчки и отчети\nв проекта Demetra', 'Title')
p('Проверка на съществуващия код и реализация на три функционалности',True)
p('Документът показва къде са реализирани филтрирането на поръчки, месечните отчети и PDF експортирането. За всяка задача са описани потребителският път, методите, файловете и резултатите от проверките. И трите функционалности липсваха в предоставения проект и са добавени в отделни основни Java файлове.')
table(['Задача','Преди промяната','Добавена реализация'],[
 ('1 Филтриране и сортиране','Само регистрации, подредени по дата за конкретен потребител','OrderFilterService.java и страница /orders'),
 ('2 Модул Отчети','Няма модул, календар или месечни обобщения','ReportService.java и страница /reports'),
 ('3 PDF преглед и изтегляне','Няма генериране или изтегляне на отчети','ReportPdfService.java и /reports/pdf'),
],[3.6,5.4,8.4])
sub('Обхват на реализацията')
p('Проектът е Spring Boot с Java 17, Spring Security, JPA и Thymeleaf. Съществуващите модели са Event, Registration и User. Под „поръчка“ тук се разбира регистрация за събитие. Страниците за поръчки и отчети са на ниво ORGANIZER и показват само неговите данни.')
p('Добавен е модел Payment за ръчно отбелязани получени плащания. Така събраните пари се изчисляват от записи за плащане, а неплатените регистрации не се приемат автоматично за приход. Валутата е USD според съществуващите цени с $ в проекта.')
sub('Как са получени изображенията')
p('Проверката използва реалните контролери и Thymeleaf шаблони през Spring MockMvc с отделна H2 база и синтетични данни. HTML изображенията са печатни визуализации през LibreOffice с опростено оформление; те не са скрийншотове от интерактивен браузър. Изображенията на PDF са от действително генерирания файл. Кодовите изображения са точни извадки с номера на редовете.')
p('Docker достъпът и отварянето на локален сървърен порт са блокирани в средата. Не е стартиран контейнер и съществуващият Demetra не е променян. Подготвена и валидирана е отделна Compose конфигурация. Всичките 16 автоматизирани теста преминават успешно.')

# 2
page(); h('Задача 1 Филтриране и сортиране на поръчки')
p('Потребителски път: вход като организатор → „Поръчки“ в навигацията → избор на филтри → „Филтрирай“. „Изчисти“ връща всички регистрации за събитията на текущия организатор.')
code('GET /orders → OrderController.list → OrderFilterService.filterAndSort\n→ RegistrationRepository.findByEventOrganizerId → orders/list.html')
p('Основният метод е в src/main/java/org/codingburgas/demetra/services/OrderFilterService.java, ред 27. Контролерът е OrderController.java, ред 14; интерфейсът е src/main/resources/templates/orders/list.html.')
p('Седмицата е по ISO и започва в понеделник. Точната дата, денят от седмицата и часът се проверяват по Registration.createdAt. Филтрите се комбинират с И. Час 9 включва целия интервал 09:00–09:59. Подреждането е по дата и час, а при равенство — по ID; има възходящ и низходящ ред.')
picture('code-orders.png','Фигура 1 Реалният метод filterAndSort с проверка на параметрите, комбинирани филтри и сортиране.',6.7)
p('Невалиден ден, час, ISO седмица или посока на сортиране връща HTTP 400. Repository заявката първо ограничава записите до събитията на текущия организатор.')

# 3
page(); h('Задача 1 Пример с комбинирани филтри')
p('При седмица 2026-W42, дата 16.10.2026, ден петък, час 9 и възходящ ред остава точно една поръчка — регистрацията на Анна Иванова за „Есенен концерт“ в 09:15. Сумата 25.00 USD е вече отбелязано получено плащане.')
code('/orders?week=2026-W42&date=2026-10-16&day=5&hour=9&sort=asc')
picture('orders-print-crop.png','Фигура 2 Печатен изглед на действителния HTML отговор за комбинирания филтър. Оформлението е адаптирано за печат; данните и избраните стойности са от приложението.',6.0)
p('Проверка: ReportsEvidenceTests потвърждава „Намерени поръчки: 1“. ReportsFlowTests отделно проверява ISO седмица, която преминава през границата на годината, граничните минути на часа и двете посоки на сортиране.')

# 4
page(); h('Задача 2 Модул Отчети')
p('Потребителски път: вход като организатор → отделния бутон „Отчети“ → дата от крайния месец → избор 1, 6 или 12 месеца → „Покажи отчет“. Календарът е HTML input с type="date"; периодът е списък с трите допустими избора.')
code('GET /reports → ReportController.index → ReportService.generate\n→ EventRepository + RegistrationRepository + PaymentRepository\n→ reports/index.html')
p('Основният метод е в services/ReportService.java, ред 30, под общия Java път от задача 1. Контролерът е ReportController.java, ред 17; шаблонът е src/main/resources/templates/reports/index.html. Навигацията е във fragments/nav.html.')
p('Избраната дата определя крайния месец. Началото е първият ден на най-ранния месец, а краят за заявката е първият ден на следващия месец, изключително. Целият период трябва да попада в последните 12 календарни месеца, включително текущия. За текущия месец стойностите са междинни.')
picture('code-reports.png','Фигура 3 Реализация на календарните граници, ограничението до 12 месеца и обединяването на събития, регистрации и плащания.',6.7)
p('Събитията се избират по Event.startsAt. Показват се всички статуси и категории; събития без дата не участват. За всеки месец, включително празните, се връщат брой събития, потвърдени регистрации и сума на получените плащания.')

# 5
page(); h('Задача 2 Календар и резултати')
picture('reports-controls-crop.png','Фигура 4 Печатна визуализация на HTML отговора за 1 месец. Виждат се календарното поле, периодът, двата PDF линка и обобщението. Това е статичен изглед с оформление за печат.',6.3)
p('Резултати за избрана дата 16.10.2026, преди допълнителното плащане в теста:')
table(['Период','Календарни граници','Събития','Потвърдени','Получени USD'],[
 ('1 месец','01.10.2026 – 31.10.2026',2,3,'65.00'),
 ('6 месеца','01.05.2026 – 31.10.2026',5,9,'140.00'),
 ('12 месеца','01.11.2025 – 31.10.2026',6,11,'165.00'),
],[2.0,6.1,2.3,3.0,4.0])
p('Сумите са всички ръчно отбелязани постъпления за избраните събития, независимо от датата на плащане. Това е справка по дата на събитие, а не касов дневник по дата на постъпване. Няма частични плащания или възстановявания. Отказана регистрация не заличава вече записана разписка.')

# 6
page(); h('Задача 3 Преглед и изтегляне на PDF')
p('От страницата „Отчети“ бутонът „Преглед на PDF“ отваря файла в нов раздел, а „Изтегли PDF“ задава изтегляне. И двата използват същите date и months, с които е генериран HTML отчетът.')
code('GET /reports/pdf → ReportController.export → ReportService.generate\n→ ReportPdfService.generate → ResponseEntity<byte[]>')
p('Основният файл е services/ReportPdfService.java, метод generate на ред 14. ReportController.export на ред 32 връща application/pdf и избира inline или attachment според параметъра download.')
picture('code-pdf-route.png','Фигура 5 Контролерът предава същия период, генерира PDF и задава HTTP заглавката за преглед или изтегляне.',6.7)
p('PDF включва общите показатели, всеки месечен ред и всички събития с категория, дата, статус, регистрации и получени суми. Използва Apache PDFBox и вграден DejaVu Sans за кирилица. Дългите текстове се пренасят, а съдържанието продължава на следващи страници.')
p('Провереният отговор за изтегляне е:')
code('Status: 200\nContent-Type: application/pdf\nContent-Disposition: attachment;\n filename="demetra-report-2025-11-01-2026-10-31.pdf"')
picture('pdf-events-crop.png','Фигура 6 Извадка от втората страница на реално генерирания PDF. Показани са останалите две събития и техните суми.',6.5)
p('Модулът експортира всички добавени справки за събития. Проектът не съдържа пълно счетоводство, фактури, данъчни справки или разходни регистри.')

# 7
page(); h('Задача 3 Генерираният PDF отчет')
picture('report-12-1.png','Фигура 7 Първа страница от действително генерирания 12-месечен отчет. Втората страница продължава списъка със събития. Данните са синтетични.',5.65)

# 8
page(); h('Проверки на реализацията')
p('Изпълнена команда: mvn -o test. Резултат на 16.10.2026: 16 теста, 0 неуспешни проверки, 0 грешки и 0 пропуснати теста. Седем от тестовете са съществуващи; осем са нови функционални проверки и един генерира доказателствените HTML и PDF файлове.')
table(['Проверка','Потвърден резултат'],[
 ('Седмица, дата, ден и час','Комбинираният филтър връща само съвпадащите записи; проверена е и ISO граница между години.'),
 ('Сортиране','Възходящ и низходящ ред по дата и час; ID решава равенствата.'),
 ('Входни параметри','Отхвърлят се невалидни седмици, дни, часове, посоки и периоди.'),
 ('Календарни граници','Началният момент участва; първият момент на следващия месец не участва. Има редове и за празни месеци.'),
 ('Плащания','Повторното отбелязване не удвоява сумата; промяна на цената след плащане не променя записаната сума.'),
 ('Права за достъп','Студентите получават 403; анонимните се пренасочват към вход. Чужди събития и плащания са недостъпни.'),
 ('Форми и HTML','Рендират се реалните Thymeleaf шаблони; има CSRF защита при отбелязване на плащане.'),
 ('PDF','Проверени са кирилица, суми, inline/attachment заглавки, дълги заглавия и всички 24 тестови събития на няколко страници.'),
],[4.0,13.4])
p('Тестовите файлове са src/test/java/org/codingburgas/demetra/ReportsFlowTests.java и ReportsEvidenceTests.java. Резултатите са в target/surefire-reports/. Повторното mvn test генерира HTML и примерен PDF в target/report-evidence/.')
p('Изпълнено е и успешно пакетиране: mvn -o -Dmaven-jar-plugin.version=3.4.1 -DskipTests package. Получен е target/Demetra-0.0.1-SNAPSHOT.jar. Параметърът използва наличния локален кеш; версията в pom.xml не е сменяна.')
p('Ограничение на проверката: изпълнението е през MockMvc и H2. Не е проведена проверка в работещ Docker/MySQL контейнер или интерактивна проверка на браузърните форми. Проверена е синтактичната валидност на отделната Compose конфигурация.')

# 9
page(); h('Стартиране и ориентация във файловете')
sub('Отделен Docker проект')
code('docker compose -f compose.reports.yml up -d --build')
p('Файлът е самостоятелен. Име на проекта: demetra-dariq-reports. Контейнери: demetra-dariq-reports-app-1 и demetra-dariq-reports-db-1. Приложението е на http://localhost:8093. Базата няма публикуван хост порт и използва собствен том demetra-dariq-reports_reports_data. Оригиналният docker-compose.yml е запазен.')
p('Compose вариантът стартира с празна база. Създайте Teacher профил, добавете събития и регистрации и отбележете получените плащания. Конфигурацията е подготвена и проверена с docker compose -f compose.reports.yml config --quiet; контейнерите не са стартирани в тази среда.')
sub('Локална демонстрация')
code('mvn spring-boot:run -Dspring-boot.run.profiles=demo')
p('Адрес: http://localhost:8093. Демонстрационен профил: reports@example.com; парола: ReportsDemo123. Данните са синтетични и са в отделна временна H2 база. При спиране се губят. Профилът demo не трябва да се включва за реална база.')
sub('Основни добавени файлове')
table(['Път спрямо проекта','Предназначение'],[
 ('src/main/java/org/codingburgas/demetra/services/\nOrderFilterService.java','Задача 1 и отбелязване на плащане'),
 ('src/main/java/org/codingburgas/demetra/services/\nReportService.java','Задача 2 и общ модел на отчета'),
 ('src/main/java/org/codingburgas/demetra/services/\nReportPdfService.java','Задача 3 и многостраничен PDF'),
 ('src/main/java/org/codingburgas/demetra/controllers/\nOrderController.java и ReportController.java','HTTP маршрути за страниците и PDF'),
 ('src/main/resources/templates/\norders/list.html и reports/index.html','Форми, таблици, календар и PDF бутони'),
 ('src/main/java/org/codingburgas/demetra/models/Payment.java\nsrc/main/java/org/codingburgas/demetra/repositories/PaymentRepository.java','Записване и извличане на получените суми'),
 ('compose.reports.yml и application-demo.properties','Отделно Docker и демонстрационно стартиране'),
],[11.8,5.6])
p('Допълнени са SecurityConfig, навигацията, EventRepository, RegistrationRepository, CSS и pom.xml. Новите зависимости са PDFBox и H2 при изпълнение. В src/main/resources/fonts/ са шрифтът и лицензът му. Подробни инструкции има и в REPORTS_README.md.')

# Ensure table rows do not split between pages.
for t in doc.tables:
    for row in t.rows:
        no=OxmlElement('w:cantSplit'); row._tr.get_or_add_trPr().append(no)
for border in doc.styles.element.xpath('.//w:pBdr'):
    border.getparent().remove(border)
for border in doc.element.xpath('.//w:pBdr'):
    border.getparent().remove(border)

doc.core_properties.title='Поръчки и отчети в проекта Demetra'
doc.core_properties.subject='Реализация и проверка на три функционалности'
doc.core_properties.author=''
doc.core_properties.keywords='Demetra, поръчки, отчети, PDF'
file=OUT/'Demetra_Три_задачи_описание.docx'
doc.save(file)
print(file)
